"""Builds the input files for the real-file check (PLAN.md stage 4).

Usage: python3 make-files.py OUT_DIR   (needs requirements.txt: python-docx, openpyxl, python-pptx, Pillow)

Each Office type is built twice: `full-*` carries every kind of mark, `clean-*` is an ordinary Russian document with
none. Bases come from the Office libraries, so they are valid packages; marks are then written into them. Next to each
input a `*.expected.json` lists the text that must survive cleaning, so a run can be checked without reading by eye.

Text is written with marked spans: M("x") wraps a carrier that must disappear. dirty() keeps it, clean() drops it.
"""

from __future__ import annotations

import io
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table
from PIL import Image
from PIL.PngImagePlugin import PngInfo
from pptx import Presentation
from pptx.util import Inches

ZW, ZWNJ, ZWJ, WJ, BOM = "​", "‌", "‍", "⁠", "﻿"
TAGS = "\U000e0041\U000e0042"
VS = "\U000e0100"
NBSP, NNBSP, SHY, LRM = " ", " ", "­", "‎"
FAMILY = "\U0001f468‍\U0001f469‍\U0001f467"
FLAG = "\U0001f3f4" + "".join(chr(0xE0000 + ord(c)) for c in "gbeng") + "\U000e007f"
HEART = "❤️"

OPEN, CLOSE = "", ""


def M(s: str) -> str:
    return OPEN + s + CLOSE


def dirty(s: str) -> str:
    return s.replace(OPEN, "").replace(CLOSE, "")


def clean(s: str) -> str:
    return re.sub(OPEN + ".*?" + CLOSE, "", s, flags=re.S)


# Legitimate typography and scripts must survive with the default options; the marked spans must not.
PARAGRAPHS = [
    "При" + M(ZW) + "вет, мир! Это тест" + M(WJ) + "овый абзац.",
    "«Ёлочки»" + NBSP + "—" + NBSP + "тире, цена 1" + NNBSP + "000" + NBSP + "₽, пере" + M(SHY) + "нос.",  # upstream removes U+00AD (docs/known-issues.md)
    "Беларуская: ў, і, ґ. Lietuvių: ąčęėįšųūž. Українська: ґєії.",
    "Emoji: " + FAMILY + " " + FLAG + " " + HEART + " ok" + M(ZWNJ + ZWJ) + ".",
    "Tags" + M(TAGS) + " and" + M(VS) + " variation" + M(BOM) + " selectors.",
    "Direction marks stay: abc" + LRM + " def (strip_bidi is off).",
    "Look-alike letters stay by default: Аpple (Cyrillic А).",
    "Tom & Jerry <b> \"q\" 'x' ]]> " + M(ZW) + "end.",
]
SHORT = [PARAGRAPHS[0], PARAGRAPHS[1], PARAGRAPHS[3], PARAGRAPHS[7]]


def runs(s: str, marked: bool) -> list[str]:
    """Splits a text into runs so that each marked span sits in a run of its own (Word does this in real files)."""
    if not marked:
        return [clean(s)]
    return [p[1:-1] if p.startswith(OPEN) else p for p in re.split("(" + OPEN + ".*?" + CLOSE + ")", s, flags=re.S) if p]


def write_runs(paragraph, s: str, marked: bool) -> None:
    for text in runs(s, marked):
        paragraph.add_run().text = text


def png(marked: bool) -> bytes:
    info = PngInfo()
    if marked:
        info.add_text("Software", "DALL-E 3")
        info.add_text("Comment", "AI generated image")
    buf = io.BytesIO()
    Image.new("RGB", (32, 32), (200, 30, 30)).save(buf, format="PNG", pnginfo=info)
    return buf.getvalue()


def jpeg(marked: bool) -> bytes:
    exif = Image.Exif()
    if marked:
        exif[0x010F] = "OpenAI"
        exif[0x0131] = "DALL-E 3"
    buf = io.BytesIO()
    Image.new("RGB", (32, 32), (30, 30, 200)).save(buf, format="JPEG", exif=exif)
    return buf.getvalue()


# ---- Writing marks into a finished package ---------------------------------------------------------------------

REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
APP_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
    "<Application>ChatGPT</Application><AppVersion>16.0000</AppVersion><Company>OpenAI</Company>"
    "<Manager>GPT</Manager></Properties>"
)
CUSTOM_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="ai_generator">'
    "<vt:lpwstr>ChatGPT</vt:lpwstr></property></Properties>"
)
ITEM_XML = '<?xml version="1.0" encoding="UTF-8"?><ai xmlns="urn:example:provenance"><generator>ChatGPT</generator></ai>'
ITEM_PROPS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<ds:datastoreItem ds:itemID="{6F2A6A2E-1111-4B7B-9A3A-0123456789AB}" '
    'xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml"><ds:schemaRefs/></ds:datastoreItem>'
)
WEBEXT = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<we:webextension xmlns:we="http://schemas.microsoft.com/office/webextensions/webextension/2010/11" '
    'id="{11111111-2222-3333-4444-555555555555}"><we:reference id="WA200000000" version="1.0.0.0" store="OMEX" '
    'storeType="OMEX"/><we:alternateReferences/><we:properties/><we:bindings/><we:snapshot/></we:webextension>'
)
TASKPANES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<wetp:taskpanes xmlns:wetp="http://schemas.microsoft.com/office/webextensions/taskpanes/2010/11">'
    '<wetp:taskpane dockstate="right" visibility="0" width="350" row="4">'
    '<wetp:webextensionref xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/>'
    "</wetp:taskpane></wetp:taskpanes>"
)


def add_rel(rels: str, rid: str, rel_type: str, target: str) -> str:
    return rels.replace("</Relationships>", f'<Relationship Id="{rid}" Type="{rel_type}" Target="{target}"/></Relationships>')


def add_override(ct: str, part: str, ctype: str) -> str:
    if f'PartName="{part}"' in ct:
        return ct
    return ct.replace("</Types>", f'<Override PartName="{part}" ContentType="{ctype}"/></Types>')


def write_package(path: Path, parts: dict[str, bytes], order: list[str]) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name in order:
            z.writestr(name, parts[name])


def inject_marks(path: Path, main_rels: str, webextension: bool) -> None:
    """Adds provenance metadata, a customXml tree (and for DOCX a task-pane add-in) to a package built by a library."""
    with zipfile.ZipFile(path) as z:
        order = z.namelist()
        parts = {n: z.read(n) for n in order}
    ct = parts["[Content_Types].xml"].decode()
    root = parts["_rels/.rels"].decode()
    rels = parts[main_rels].decode()
    rels_dir = main_rels.rsplit("_rels/", 1)[0]

    for name, body, rel_type, ctype in [
        ("docProps/app.xml", APP_XML, REL + "/extended-properties", "application/vnd.openxmlformats-officedocument.extended-properties+xml"),
        ("docProps/custom.xml", CUSTOM_XML, REL + "/custom-properties", "application/vnd.openxmlformats-officedocument.custom-properties+xml"),
    ]:
        if name not in parts:
            order.append(name)
            if rel_type not in root:
                root = add_rel(root, "rIdMark" + name[-7:-4], rel_type, name)
        parts[name] = body.encode()
        ct = add_override(ct, "/" + name, ctype)

    for name, body in [
        ("customXml/item1.xml", ITEM_XML),
        ("customXml/itemProps1.xml", ITEM_PROPS),
        ("customXml/_rels/item1.xml.rels", add_rel('<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>', "rId1", REL + "/customXmlProps", "itemProps1.xml")),
    ]:
        if name not in parts:
            order.append(name)
        parts[name] = body.encode()
    ct = add_override(ct, "/customXml/itemProps1.xml", "application/vnd.openxmlformats-officedocument.customXmlProperties+xml")
    if "customXml/item1.xml" not in rels:  # the python-docx template already ships one
        rels = add_rel(rels, "rIdMarkXml", REL + "/customXml", "../customXml/item1.xml")

    if webextension:
        for name, body in [
            ("word/webextensions/webextension1.xml", WEBEXT),
            ("word/webextensions/taskpanes.xml", TASKPANES),
            ("word/webextensions/_rels/taskpanes.xml.rels", add_rel('<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>', "rId1", "http://schemas.microsoft.com/office/2011/relationships/webextension", "webextension1.xml")),
        ]:
            order.append(name)
            parts[name] = body.encode()
        ct = add_override(ct, "/word/webextensions/webextension1.xml", "application/vnd.ms-office.webextension+xml")
        ct = add_override(ct, "/word/webextensions/taskpanes.xml", "application/vnd.ms-office.webextensiontaskpanes+xml")
        root = add_rel(root, "rIdMarkWeb", "http://schemas.microsoft.com/office/2011/relationships/webextensiontaskpanes", "word/webextensions/taskpanes.xml")

    parts["[Content_Types].xml"], parts["_rels/.rels"], parts[main_rels] = ct.encode(), root.encode(), rels.encode()
    write_package(path, parts, order)
    assert rels_dir  # main_rels always lives in a folder


# ---- Builders --------------------------------------------------------------------------------------------------


def set_props(props, marked: bool) -> None:
    props.title = "Отчёт"
    props.author = "ChatGPT" if marked else "Иван Петров"
    props.last_modified_by = "OpenAI GPT" if marked else "Иван Петров"
    props.comments = "Generated by ChatGPT" if marked else "Квартальный отчёт"
    props.keywords = "ai, chatgpt" if marked else "отчёт"
    props.subject = "AI draft" if marked else "Отчёт"
    props.category = "AI" if marked else "Работа"


def build_docx(path: Path, marked: bool) -> list[str]:
    d = Document()
    set_props(d.core_properties, marked)
    for s in PARAGRAPHS:
        write_runs(d.add_paragraph(), s, marked)
    table = d.add_table(rows=2, cols=2)
    for cell, s in zip((c for r in table.rows for c in r.cells), SHORT):
        write_runs(cell.paragraphs[0], s, marked)
    write_runs(d.sections[0].header.paragraphs[0], PARAGRAPHS[0], marked)
    first = d.add_paragraph()
    write_runs(first, "Комментируемый" + M(ZW) + " текст.", marked)
    d.add_comment(runs=first.runs[0], text=dirty("Проверить" + M(ZW) + " это") if marked else "Проверить это", author="ChatGPT" if marked else "Иван", initials="CG")
    d.add_picture(io.BytesIO(png(marked)))
    d.add_picture(io.BytesIO(jpeg(marked)))
    d.save(path)
    if marked:
        inject_marks(path, "word/_rels/document.xml.rels", webextension=True)
    return [clean(s) for s in PARAGRAPHS + SHORT] + ["Комментируемый текст.", "Проверить это"]


def build_xlsx(path: Path, marked: bool) -> list[str]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    rows = [
        ("Имя" + M(ZW), "Сумма"),
        ("Иван" + NBSP + "Петров", 100),
        (PARAGRAPHS[0], 250.5),
        (PARAGRAPHS[1], 0),
    ]
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, clean(v) if isinstance(v, str) and not marked else dirty(v) if isinstance(v, str) else v)
    ws["D1"] = "=SUM(B2:B4)"
    ws["D2"] = '=CONCATENATE("при' + (ZW if marked else "") + 'вет"," ","мир")'
    ws.add_table(Table(displayName="Tbl", ref="A1:B4"))
    ws["A2"].comment = Comment(dirty("Из " + M(ZW) + "ChatGPT") if marked else "Проверено", "ChatGPT" if marked else "Иван")
    ws2 = wb.create_sheet("Лист2")
    ws2["A1"] = dirty(PARAGRAPHS[3]) if marked else clean(PARAGRAPHS[3])
    p = wb.properties
    p.title, p.creator, p.lastModifiedBy = "Отчёт", "ChatGPT" if marked else "Иван Петров", "OpenAI GPT" if marked else "Иван Петров"
    p.description, p.keywords = "Generated by ChatGPT" if marked else "Отчёт", "ai"
    wb.save(path)
    if marked:
        inject_marks(path, "xl/_rels/workbook.xml.rels", webextension=False)
    return [clean(v) for r in rows for v in r if isinstance(v, str)] + [clean(PARAGRAPHS[3])]


def build_pptx(path: Path, marked: bool) -> list[str]:
    prs = Presentation()
    set_props(prs.core_properties, marked)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    write_runs(slide.shapes.title.text_frame.paragraphs[0], PARAGRAPHS[0], marked)
    body = slide.placeholders[1].text_frame
    write_runs(body.paragraphs[0], PARAGRAPHS[1], marked)
    write_runs(body.add_paragraph(), PARAGRAPHS[3], marked)
    notes = slide.notes_slide.notes_text_frame
    write_runs(notes.paragraphs[0], "Заметки" + M(ZW) + " докладчика", marked)
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    write_runs(slide2.shapes.title.text_frame.paragraphs[0], PARAGRAPHS[7], marked)
    slide2.shapes.add_picture(io.BytesIO(png(marked)), Inches(1), Inches(2))
    prs.save(path)
    if marked:
        inject_marks(path, "ppt/_rels/presentation.xml.rels", webextension=False)
    return [clean(PARAGRAPHS[0]), clean(PARAGRAPHS[1]), clean(PARAGRAPHS[3]), "Заметки докладчика", clean(PARAGRAPHS[7])]


# ---- Main ------------------------------------------------------------------------------------------------------


def main(out: Path) -> None:
    (out / "text").mkdir(parents=True, exist_ok=True)
    (out / "office").mkdir(exist_ok=True)
    (out / "errors").mkdir(exist_ok=True)

    # Text: a Windows file (BOM, CRLF), a pasted text (LF), and a legacy encoding that is not UTF-8.
    (out / "text/ru-mixed.txt").write_bytes((BOM + "\r\n".join(dirty(s) for s in PARAGRAPHS) + "\r\n").encode())
    (out / "text/ru-mixed-paste.txt").write_text("\n".join(dirty(s) for s in PARAGRAPHS), encoding="utf-8")
    (out / "text/ru-mixed.expected.json").write_text(json.dumps([clean(s) for s in PARAGRAPHS], ensure_ascii=False, indent=1))
    (out / "text/cp1251.txt").write_bytes("Привет, мир".encode("cp1251"))

    for kind, build in (("docx", build_docx), ("xlsx", build_xlsx), ("pptx", build_pptx)):
        for marked in (True, False):
            name = f"{'full' if marked else 'clean'}-{kind}"
            expected = build(out / "office" / f"{name}.{kind}", marked)
            (out / "office" / f"{name}.expected.json").write_text(json.dumps(expected, ensure_ascii=False, indent=1))

    # Inputs the page must refuse with a message instead of failing or reporting success.
    docx = (out / "office/full-docx.docx").read_bytes()
    (out / "errors/truncated.docx").write_bytes(docx[: len(docx) * 6 // 10])
    (out / "errors/text-as-docx.docx").write_text("Привет, мир", encoding="utf-8")
    (out / "errors/encrypted.docx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504 + "EncryptedPackage".encode("utf-16-le"))
    (out / "errors/legacy.doc").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504)
    (out / "errors/empty.docx").write_bytes(b"")
    (out / "errors/macro.docm").write_bytes(docx)


if __name__ == "__main__":
    main(Path(sys.argv[1]))
